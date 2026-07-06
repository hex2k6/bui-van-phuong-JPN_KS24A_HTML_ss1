import os
import base64
import json
import requests
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# --- CONFIG & DIRECTORY SETUP ---
OUTPUT_DIR = r"c:\Users\D365\Documents\BTVN\AI\SS13\Smart_Hospital_SRS"
os.makedirs(OUTPUT_DIR, exist_ok=True)

PROMPTS_FILE = os.path.join(OUTPUT_DIR, "History_Prompts.txt")
DOCX_FILE = os.path.join(OUTPUT_DIR, "System_Design_SRS.docx")

# --- COLOR SYSTEM (Modern Blue theme) ---
COLOR_PRIMARY = (16, 44, 87)       # Deep Blue #102C57
COLOR_SECONDARY = (53, 114, 239)   # Medium Blue #3572EF
COLOR_TEXT = (43, 43, 43)          # Charcoal #2B2B2B
COLOR_MUTED = (108, 117, 125)      # Slate Gray #6C757D
HEX_BG_HEADER = "102C57"           # Hex for cell shading
HEX_BG_ZEBRA = "F2F7FF"            # Zebra striping light blue

# --- FUNCTIONS FOR WORD FORMATTING ---
def set_cell_background(cell, hex_color):
    """Set background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set padding for cell in twentieths of a point (dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for margin, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(margin)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_borders(cell, color="CCCCCC", sz="4", val="single"):
    """Apply borders to a cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for border_name in ['top', 'left', 'bottom', 'right']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), val)
        border.set(qn('w:sz'), sz)
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), color)
        tcBorders.append(border)
    tcPr.append(tcBorders)

def add_styled_heading(doc, text, level):
    """Add a heading with consistent styling and hierarchy."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    
    run = p.add_run(text)
    if level == 1:
        run.font.name = "Arial"
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor(*COLOR_PRIMARY)
        run.bold = True
        # Add left accent border or divider below if possible, here we use simple layout
    elif level == 2:
        run.font.name = "Arial"
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(*COLOR_SECONDARY)
        run.bold = True
    elif level == 3:
        run.font.name = "Arial"
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(*COLOR_PRIMARY)
        run.bold = True
        run.italic = True
    return p

def add_body_paragraph(doc, text="", bold=False, space_after=6):
    """Add body paragraph with proper font, size, line spacing and spacing after."""
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(space_after)
    if text:
        run = p.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(*COLOR_TEXT)
        run.bold = bold
    return p

def add_bullet_point(doc, text):
    """Add a customized bullet point paragraph."""
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor(*COLOR_TEXT)
    return p

def create_styled_table(doc, rows, cols, headers, col_widths=None):
    """Create a table with headers and consistent styling."""
    table = doc.add_table(rows=rows, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Format headers
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], HEX_BG_HEADER)
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=150, right=150)
        set_cell_borders(hdr_cells[i], color="102C57", sz="6")
        
        # Center header text or left align depending on preference
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.runs[0]
        run.font.name = "Arial"
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.bold = True
        
        if col_widths and i < len(col_widths):
            hdr_cells[i].width = Inches(col_widths[i])
            
    return table

def populate_table_row(table, row_idx, data, col_widths=None, zebra=False):
    """Populate data into a table row and apply styling."""
    row_cells = table.rows[row_idx].cells
    bg_color = HEX_BG_ZEBRA if zebra else "FFFFFF"
    for i, val in enumerate(data):
        row_cells[i].text = str(val)
        set_cell_background(row_cells[i], bg_color)
        set_cell_margins(row_cells[i], top=100, bottom=100, left=150, right=150)
        set_cell_borders(row_cells[i], color="E0E0E0", sz="4")
        
        p = row_cells[i].paragraphs[0]
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(2)
        if p.runs:
            run = p.runs[0]
            run.font.name = "Arial"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(*COLOR_TEXT)
            
        if col_widths and i < len(col_widths):
            row_cells[i].width = Inches(col_widths[i])

# --- MERMAID RENDERING ---
def fetch_mermaid_image(mermaid_code, filename):
    """Fetch PNG image of a Mermaid diagram using mermaid.ink."""
    print(f"Fetching Mermaid image for {filename}...")
    data = {
        "code": mermaid_code,
        "mermaid": {"theme": "neutral"}
    }
    json_str = json.dumps(data)
    b64_bytes = base64.urlsafe_b64encode(json_str.encode('utf-8'))
    b64_str = b64_bytes.decode('utf-8').rstrip('=')
    
    url = f"https://mermaid.ink/img/{b64_str}"
    filepath = os.path.join(OUTPUT_DIR, filename)
    try:
        response = requests.get(url, timeout=15)
        if response.status_code == 200:
            with open(filepath, "wb") as f:
                f.write(response.content)
            print(f"Successfully saved diagram to {filepath}")
            return filepath
        else:
            print(f"Mermaid.ink returned status {response.status_code}: {response.text}")
    except Exception as e:
        print(f"Error fetching Mermaid image: {e}")
    return None

# --- GENERATE HISTORY_PROMPTS.TXT ---
def generate_prompts_file():
    """Create the History_Prompts.txt file containing the prompt sequence."""
    print("Generating History_Prompts.txt...")
    prompts = [
        "========================================================================\n"
        "LỊCH SỬ CÁC CÂU LỆNH PROMPT - THIẾT KẾ HỆ THỐNG SMART HOSPITAL (SRS)\n"
        "========================================================================\n\n",
        
        "[PROMPT 1 - Khởi tạo dự án & Định hình bối cảnh]\n"
        "\"Chào bạn, tôi là một chuyên gia phân tích nghiệp vụ và kiến trúc sư hệ thống độc lập. "
        "Tôi cần bạn đóng vai trò là một AI Assistant cao cấp để cùng thiết kế một Bản đặc tả yêu cầu "
        "phần mềm (SRS) cho dự án 'Hệ thống Quản lý Bệnh viện Đa khoa Thông minh' (Smart Hospital). "
        "Yêu cầu ban đầu là hệ thống phải bao gồm ít nhất 3 phân hệ cốt lõi và có sự tham gia của 3 đối tượng "
        "người dùng khác nhau. Hãy liệt kê đề xuất của bạn về các Actors và Modules chính kèm mô tả ngắn gọn "
        "để chúng ta thống nhất trước khi đi vào chi tiết.\"\n\n",
        
        "[PROMPT 2 - Làm sâu sắc vai trò và đặc tả các Modules]\n"
        "\"Bản đề xuất rất tốt. Chúng ta sẽ chốt phương án:\n"
        "- 3 Actors: Bệnh nhân (Patient), Bác sĩ (Doctor), và Nhân viên Y tế/Admin (Medical Staff/Admin).\n"
        "- 3 Phân hệ: Phân hệ Đăng ký & Lịch hẹn, Phân hệ Hồ sơ bệnh án điện tử (EHR), Phân hệ Viện phí & Thanh toán.\n"
        "Bây giờ, hãy viết chi tiết hơn về các chức năng chính của từng phân hệ dưới dạng các User Stories "
        "cho từng đối tượng. Đối với mỗi User Story, hãy chỉ rõ: Story ID, Đối tượng tác động (Actor), "
        "Nội dung Story (As a... I want to... So that...), và Tiêu chí chấp nhận (Acceptance Criteria). "
        "Vui lòng trình bày dưới dạng bảng cho rõ ràng.\"\n\n",
        
        "[PROMPT 3 - Thiết kế Sơ đồ Use Case tổng quan]\n"
        "\"Hãy thiết kế một sơ đồ Use Case tổng quan cho hệ thống này bằng ngôn ngữ mô tả Mermaid. "
        "Sơ đồ cần phân chia rõ ràng các phân hệ bằng subgraph và thể hiện mối quan hệ liên kết giữa các "
        "tác nhân (Patient, Doctor, Admin) với các use cases tương ứng của họ. Đảm bảo cú pháp Mermaid "
        "chuẩn xác để có thể render được trực tiếp.\"\n\n",
        
        "[PROMPT 4 - Phân tích & Đặc tả Phân hệ Viện phí (Bổ sung trạng thái)]\n"
        "\"Trong Phân hệ Viện phí & Thanh toán, quy trình xử lý hóa đơn cần rõ ràng hơn. Hãy định nghĩa "
        "rõ các trạng thái của hóa đơn (Ví dụ: UNPAID, PAID, CANCELLED) và các hình thức thanh toán được chấp nhận "
        "(CASH, CREDIT_CARD, BANK_TRANSFER, E_WALLET). Đồng thời, bổ sung yêu cầu phi chức năng về bảo mật thông tin "
        "thanh toán và tuân thủ các tiêu chuẩn an toàn dữ liệu y tế (HIPAA).\"\n\n",
        
        "[PROMPT 5 - Thiết kế Cơ sở dữ liệu & Sơ đồ ERD]\n"
        "\"Bây giờ chúng ta sẽ chuyển sang phần thiết kế Cơ sở dữ liệu. Hãy xây dựng một sơ đồ ERD bằng Mermaid "
        "thể hiện các bảng dữ liệu cốt lõi gồm: Patients, Doctors, Appointments, MedicalRecords, Bills. "
        "Chỉ rõ các trường khóa chính (PK), khóa ngoại (FK), và các mối quan hệ (1-N, N-N, 1-1). "
        "Đồng thời, hãy viết từ điển dữ liệu (Data Dictionary) chi tiết dưới dạng bảng cho từng bảng thực thể này, "
        "gồm: Tên trường, Kiểu dữ liệu, Ràng buộc (PK, FK, Nullable...), và Mô tả chức năng.\"\n\n",
        
        "[PROMPT 6 - Bổ sung Bảo mật & Yêu cầu phi chức năng cho EHR]\n"
        "\"Hồ sơ bệnh án điện tử (EHR) là dữ liệu cực kỳ nhạy cảm. Hãy viết thêm các yêu cầu phi chức năng chi tiết "
        "cho phân hệ EHR bao gồm: Kiểm soát truy cập dựa trên vai trò (RBAC), Nhật ký lịch sử thao tác (Audit trail), "
        "Mã hóa dữ liệu tại chỗ (Encryption at rest) và khi truyền tải (Encryption in transit). Trình bày thật khoa học.\"\n\n",
        
        "[PROMPT 7 - Tổng hợp tài liệu SRS hoàn chỉnh và chuyên nghiệp]\n"
        "\"Tuyệt vời. Hãy tổng hợp toàn bộ các nội dung trên thành một bản Tài liệu Đặc tả Yêu cầu Phần mềm (SRS) "
        "hoàn chỉnh, viết bằng tiếng Việt chuyên nghiệp, ngôn từ trang trọng của một System Analyst. "
        "Cấu trúc tài liệu bao gồm:\n"
        "1. Giới thiệu chung (Mục tiêu, Phạm vi, Thuật ngữ)\n"
        "2. Mô tả tổng quan hệ thống (Kiến trúc, Actors, Modules)\n"
        "3. Đặc tả chức năng (Sơ đồ Use Case, Danh sách User Stories chi tiết dạng bảng)\n"
        "4. Thiết kế cơ sở dữ liệu (Sơ đồ ERD, Từ điển dữ liệu dạng bảng)\n"
        "5. Yêu cầu phi chức năng (Hiệu năng, Bảo mật & HIPAA, Khả năng mở rộng)\n"
        "Lưu ý: Không thêm bất kỳ câu chào hỏi, lời giới thiệu hay kết luận thừa thãi nào của AI (như 'Dưới đây là...', 'Hy vọng tài liệu này giúp ích...'), "
        "chỉ xuất ra nội dung đặc tả thuần túy để tôi xuất bản trực tiếp.\"\n"
    ]
    
    with open(PROMPTS_FILE, "w", encoding="utf-8") as f:
        f.writelines(prompts)
    print(f"Saved {PROMPTS_FILE}")

# --- GENERATE SYSTEM_DESIGN_SRS.DOCX ---
def generate_docx_file(usecase_img_path, erd_img_path):
    """Create the System_Design_SRS.docx document with structured sections, tables, and images."""
    print("Generating System_Design_SRS.docx...")
    doc = Document()
    
    # Set standard page margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
    # 1. COVER PAGE / TIÊU ĐỀ
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(120)
    title_p.paragraph_format.space_after = Pt(12)
    run_title = title_p.add_run("BẢN ĐẶC TẢ YÊU CẦU PHẦN MỀM (SRS)")
    set_font(run_title, size_pt=24, color_rgb=COLOR_PRIMARY, bold=True)
    
    subtitle_p = doc.add_paragraph()
    subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_p.paragraph_format.space_after = Pt(180)
    run_subtitle = subtitle_p.add_run("HỆ THỐNG QUẢN LÝ BỆNH VIỆN ĐA KHOA THÔNG MINH\n(SMART HOSPITAL SYSTEM)")
    set_font(run_subtitle, size_pt=16, color_rgb=COLOR_SECONDARY, bold=True)
    
    info_p = doc.add_paragraph()
    info_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info_p.paragraph_format.line_spacing = 1.3
    run_info = info_p.add_run(
        "Vai trò: IT Consultant & System Analyst\n"
        "Phiên bản: 1.0\n"
        "Ngày lập: 06/07/2026\n"
        "Trạng thái: Hoàn thiện thiết kế sơ bộ"
    )
    set_font(run_info, size_pt=11.5, color_rgb=COLOR_MUTED)
    
    doc.add_page_break()
    
    # --- SECTION 1: GIỚI THIỆU CHUNG ---
    add_styled_heading(doc, "1. GIỚI THIỆU CHUNG", level=1)
    
    add_styled_heading(doc, "1.1. Mục tiêu dự án", level=2)
    add_body_paragraph(doc, 
        "Hệ thống Quản lý Bệnh viện Đa khoa Thông minh (Smart Hospital System) được xây dựng nhằm tối ưu hóa "
        "các quy trình vận hành hành chính, lâm sàng và tài chính của bệnh viện. Hệ thống hướng tới việc chuyển đổi số "
        "toàn diện hoạt động khám chữa bệnh, giảm thiểu thời gian chờ đợi của bệnh nhân, nâng cao hiệu suất làm việc "
        "của bác sĩ và đảm bảo tính minh bạch, chính xác trong công tác quản lý tài chính viện phí."
    )
    
    add_styled_heading(doc, "1.2. Phạm vi hệ thống", level=2)
    add_body_paragraph(doc, 
        "Hệ thống tập trung giải quyết các bài toán nghiệp vụ cốt lõi tại các cơ sở khám chữa bệnh đa khoa, bao gồm:"
    )
    add_bullet_point(doc, "Quản lý đặt lịch hẹn khám trực tuyến và điều phối phòng khám lâm sàng.")
    add_bullet_point(doc, "Số hóa toàn bộ hồ sơ bệnh án (EHR), đơn thuốc và kết quả chẩn đoán hình ảnh của bệnh nhân.")
    add_bullet_point(doc, "Tự động hóa lập hóa đơn viện phí, tích hợp các cổng thanh toán điện tử và theo dõi trạng thái công nợ.")
    
    add_styled_heading(doc, "1.3. Thuật ngữ & Từ viết tắt", level=2)
    
    term_headers = ["Thuật ngữ / Từ viết tắt", "Ý nghĩa / Định nghĩa đầy đủ"]
    term_widths = [2.2, 4.3]
    term_table = create_styled_table(doc, rows=5, cols=2, headers=term_headers, col_widths=term_widths)
    
    terms_data = [
        ("SRS", "Software Requirements Specification - Tài liệu đặc tả yêu cầu phần mềm."),
        ("EHR", "Electronic Health Record - Hồ sơ bệnh án điện tử của bệnh nhân."),
        ("HIPAA", "Health Insurance Portability and Accountability Act - Đạo luật về trách nhiệm giải trình và chuyển đổi bảo hiểm y tế (tiêu chuẩn bảo mật dữ liệu y tế của Mỹ)."),
        ("RBAC", "Role-Based Access Control - Kiểm soát truy cập dựa trên vai trò người dùng."),
    ]
    for idx, (term, defn) in enumerate(terms_data, start=1):
        populate_table_row(term_table, idx, [term, defn], col_widths=term_widths, zebra=(idx % 2 == 0))
        
    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # --- SECTION 2: MÔ TẢ TỔNG QUAN HỆ THỐNG ---
    add_styled_heading(doc, "2. MÔ TẢ TỔNG QUAN HỆ THỐNG", level=1)
    
    add_styled_heading(doc, "2.1. Kiến trúc tổng quan", level=2)
    add_body_paragraph(doc, 
        "Hệ thống hoạt động theo mô hình Client-Server hiện đại, triển khai trên hạ tầng điện toán đám mây. "
        "Giao diện phía người dùng (Client Web/App) kết nối tới hệ thống máy chủ ứng dụng (Application Server) thông qua các RESTful API an toàn. "
        "Dữ liệu được lưu trữ tập trung tại hệ quản trị cơ sở dữ liệu quan hệ, đảm bảo tính nhất quán (ACID) của các giao dịch đặt lịch và thanh toán."
    )
    
    add_styled_heading(doc, "2.2. Các tác nhân (Actors)", level=2)
    add_body_paragraph(doc, "Hệ thống bao gồm 3 nhóm người dùng (Actors) tương tác trực tiếp:")
    add_bullet_point(doc, "Bệnh nhân (Patient): Là người đăng ký tài khoản, thực hiện đặt lịch hẹn khám, xem thông tin đơn thuốc cá nhân, nhận thông báo viện phí và thực hiện thanh toán trực tuyến.")
    add_bullet_point(doc, "Bác sĩ (Doctor): Là người tiếp nhận lịch khám từ hệ thống, xem lịch sử bệnh án, cập nhật kết quả chẩn đoán, kê đơn thuốc điện tử cho bệnh nhân trong các phiên khám.")
    add_bullet_point(doc, "Nhân viên Y tế / Admin (Medical Staff/Admin): Chịu trách nhiệm duyệt và phân phối lịch khám của bác sĩ, khởi tạo hóa đơn viện phí dựa trên chỉ định của bác sĩ, kiểm tra đối soát thanh toán và quản trị tài khoản hệ thống.")
    
    add_styled_heading(doc, "2.3. Các phân hệ cốt lõi (Core Modules)", level=2)
    add_bullet_point(doc, "Phân hệ Đăng ký & Lịch hẹn: Quản lý thông tin đăng ký, hiển thị lịch làm việc của bác sĩ theo thời gian thực, cho phép đặt/hủy/đổi lịch hẹn tự động và gửi thông báo nhắc lịch qua SMS/Email.")
    add_bullet_point(doc, "Phân hệ Hồ sơ bệnh án điện tử (EHR): Lưu trữ tập trung lịch sử khám, chẩn đoán bệnh, kết quả xét nghiệm, chẩn đoán hình ảnh và lịch sử kê đơn thuốc. Hỗ trợ bác sĩ tra cứu nhanh và ra quyết định điều trị lâm sàng.")
    add_bullet_point(doc, "Phân hệ Viện phí & Thanh toán: Tính toán chi phí khám, xét nghiệm, thuốc men theo đơn; tạo hóa đơn điện tử; tích hợp thanh toán qua thẻ tín dụng, chuyển khoản ngân hàng và ví điện tử; cập nhật trạng thái hóa đơn tự động.")

    doc.add_page_break()

    # --- SECTION 3: ĐẶC TẢ CHỨC NĂNG ---
    add_styled_heading(doc, "3. ĐẶC TẢ CHỨC NĂNG", level=1)
    
    add_styled_heading(doc, "3.1. Sơ đồ Use Case tổng quan", level=2)
    add_body_paragraph(doc, 
        "Dưới đây là sơ đồ Use Case tổng quát thể hiện các tương tác chức năng của Bệnh nhân, Bác sĩ và Nhân viên Y tế/Admin đối với 3 phân hệ cốt lõi của hệ thống."
    )
    
    # Insert Use Case Image
    if usecase_img_path and os.path.exists(usecase_img_path):
        img_p = doc.add_paragraph()
        img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        img_p.paragraph_format.space_before = Pt(8)
        img_p.paragraph_format.space_after = Pt(8)
        img_run = img_p.add_run()
        img_run.add_picture(usecase_img_path, width=Inches(5.8))
        
        caption_p = doc.add_paragraph()
        caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption_run = caption_p.add_run("Hình 3.1: Sơ đồ Use Case tổng quan hệ thống Smart Hospital")
        set_font(caption_run, size_pt=9.5, color_rgb=COLOR_MUTED, italic=True)
    else:
        add_body_paragraph(doc, "[LỖI: Không thể tải sơ đồ Use Case. Vui lòng kiểm tra kết nối mạng hoặc cú pháp Mermaid.]", bold=True)
        
    add_styled_heading(doc, "3.2. Danh sách User Stories chi tiết", level=2)
    add_body_paragraph(doc, "Chi tiết yêu cầu chức năng của hệ thống được đặc tả thông qua danh sách các User Stories dưới đây:")
    
    us_headers = ["Story ID", "Actor", "Nội dung User Story (As a... I want to...)", "Tiêu chí chấp nhận (Acceptance Criteria)"]
    us_widths = [1.0, 1.1, 2.4, 2.0]
    us_table = create_styled_table(doc, rows=7, cols=4, headers=us_headers, col_widths=us_widths)
    
    us_data = [
        ("US-01", "Bệnh nhân", "Đăng ký tài khoản & Đặt lịch hẹn trực tuyến chọn bác sĩ và khung giờ mong muốn.", "1. Lịch hiển thị theo thời gian thực.\n2. Gửi email xác nhận đặt lịch thành công sau 30s.\n3. Ngăn chặn đặt lịch trùng giờ."),
        ("US-02", "Bệnh nhân", "Xem hóa đơn và thanh toán viện phí trực tuyến qua ví điện tử hoặc tài khoản ngân hàng.", "1. Hiển thị chi tiết từng khoản phí.\n2. Cập nhật trạng thái hóa đơn sang 'PAID' ngay khi giao dịch thành công.\n3. Xuất biên lai điện tử PDF."),
        ("US-03", "Bác sĩ", "Xem danh sách lịch hẹn khám của mình trong ngày và tuần hiện tại.", "1. Giao diện trực quan sắp xếp theo thứ tự thời gian.\n2. Cho phép click vào bệnh nhân để xem bệnh án trước."),
        ("US-04", "Bác sĩ", "Ghi chép chẩn đoán y khoa và tạo đơn thuốc điện tử trực tiếp trong phiên khám bệnh.", "1. Tích hợp từ điển danh mục thuốc quốc gia.\n2. Tự động kiểm tra tương tác thuốc nguy hiểm.\n3. Khóa bệnh án sau khi bác sĩ ký số."),
        ("US-05", "Admin / NV", "Quản lý và điều phối lịch trực khám của các bác sĩ theo tuần.", "1. Hỗ trợ giao diện kéo thả trực quan.\n2. Tự động cảnh báo nếu bác sĩ bị phân lịch trùng hoặc quá giờ quy định."),
        ("US-06", "Admin / NV", "Khởi tạo hóa đơn viện phí tự động từ các chỉ định lâm sàng và đơn thuốc của bác sĩ.", "1. Lấy dữ liệu tự động từ EHR.\n2. Cho phép áp dụng mã miễn giảm hoặc bảo hiểm y tế.")
    ]
    for idx, row in enumerate(us_data, start=1):
        populate_table_row(us_table, idx, row, col_widths=us_widths, zebra=(idx % 2 == 0))
        
    doc.add_page_break()

    # --- SECTION 4: THIẾT KẾ CƠ SỞ DỮ LIỆU ---
    add_styled_heading(doc, "4. THIẾT KẾ CƠ SỞ DỮ LIỆU", level=1)
    
    add_styled_heading(doc, "4.1. Sơ đồ thực thể - mối quan hệ ERD", level=2)
    add_body_paragraph(doc, 
        "Cơ sở dữ liệu của hệ thống được thiết kế chuẩn hóa để đảm bảo tối ưu hóa tốc độ truy vấn và toàn vẹn dữ liệu. "
        "Dưới đây là sơ đồ ERD thể hiện mối quan hệ giữa các thực thể cốt lõi."
    )
    
    # Insert ERD Image
    if erd_img_path and os.path.exists(erd_img_path):
        img_p = doc.add_paragraph()
        img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        img_p.paragraph_format.space_before = Pt(8)
        img_p.paragraph_format.space_after = Pt(8)
        img_run = img_p.add_run()
        img_run.add_picture(erd_img_path, width=Inches(5.2))
        
        caption_p = doc.add_paragraph()
        caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption_run = caption_p.add_run("Hình 4.1: Sơ đồ thực thể - mối quan hệ ERD hệ thống Smart Hospital")
        set_font(caption_run, size_pt=9.5, color_rgb=COLOR_MUTED, italic=True)
    else:
        add_body_paragraph(doc, "[LỖI: Không thể tải sơ đồ ERD. Vui lòng kiểm tra kết nối mạng hoặc cú pháp Mermaid.]", bold=True)
        
    add_styled_heading(doc, "4.2. Từ điển dữ liệu (Data Dictionary)", level=2)
    add_body_paragraph(doc, "Dưới đây là cấu trúc chi tiết của các bảng dữ liệu cốt lõi trong cơ sở dữ liệu:")
    
    dict_headers = ["Tên trường", "Kiểu dữ liệu", "Ràng buộc", "Mô tả / Ý nghĩa"]
    dict_widths = [1.5, 1.2, 1.3, 2.5]
    
    # Bảng PATIENTS
    add_styled_heading(doc, "Bảng 4.2.1: PATIENTS (Thông tin bệnh nhân)", level=3)
    t_patients = create_styled_table(doc, rows=8, cols=4, headers=dict_headers, col_widths=dict_widths)
    pat_data = [
        ("patient_id", "INT", "PK, Auto Increment", "Mã định danh duy nhất của bệnh nhân."),
        ("full_name", "VARCHAR(100)", "NOT NULL", "Họ và tên đầy đủ của bệnh nhân."),
        ("phone_number", "VARCHAR(15)", "UNIQUE, NOT NULL", "Số điện thoại liên lạc."),
        ("date_of_birth", "DATE", "NOT NULL", "Ngày, tháng, năm sinh."),
        ("gender", "VARCHAR(10)", "NOT NULL", "Giới tính (Nam, Nữ, Khác)."),
        ("address", "VARCHAR(255)", "Nullable", "Địa chỉ cư trú."),
        ("email", "VARCHAR(100)", "UNIQUE, Nullable", "Địa chỉ thư điện tử.")
    ]
    for idx, row in enumerate(pat_data, start=1):
        populate_table_row(t_patients, idx, row, col_widths=dict_widths, zebra=(idx % 2 == 0))
        
    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    
    # Bảng DOCTORS
    add_styled_heading(doc, "Bảng 4.2.2: DOCTORS (Thông tin bác sĩ)", level=3)
    t_doctors = create_styled_table(doc, rows=7, cols=4, headers=dict_headers, col_widths=dict_widths)
    doc_data = [
        ("doctor_id", "INT", "PK, Auto Increment", "Mã định danh duy nhất của bác sĩ."),
        ("full_name", "VARCHAR(100)", "NOT NULL", "Họ và tên bác sĩ."),
        ("specialization", "VARCHAR(100)", "NOT NULL", "Chuyên khoa (Nội khoa, Ngoại khoa, Nhi khoa...)."),
        ("phone_number", "VARCHAR(15)", "NOT NULL", "Số điện thoại."),
        ("email", "VARCHAR(100)", "UNIQUE", "Thư điện tử công vụ."),
        ("schedule_details", "TEXT", "Nullable", "Lịch trực khám cố định trong tuần.")
    ]
    for idx, row in enumerate(doc_data, start=1):
        populate_table_row(t_doctors, idx, row, col_widths=dict_widths, zebra=(idx % 2 == 0))
        
    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    
    # Bảng APPOINTMENTS
    add_styled_heading(doc, "Bảng 4.2.3: APPOINTMENTS (Thông tin lịch hẹn)", level=3)
    t_appts = create_styled_table(doc, rows=7, cols=4, headers=dict_headers, col_widths=dict_widths)
    appt_data = [
        ("appointment_id", "INT", "PK, Auto Increment", "Mã định danh cuộc hẹn."),
        ("patient_id", "INT", "FK -> PATIENTS", "Mã bệnh nhân đặt lịch."),
        ("doctor_id", "INT", "FK -> DOCTORS", "Mã bác sĩ tiếp nhận khám."),
        ("appointment_date", "DATETIME", "NOT NULL", "Thời gian hẹn khám cụ thể."),
        ("reason", "TEXT", "Nullable", "Lý do khám bệnh sơ bộ."),
        ("status", "VARCHAR(20)", "NOT NULL", "Trạng thái: PENDING, CONFIRMED, COMPLETED, CANCELLED.")
    ]
    for idx, row in enumerate(appt_data, start=1):
        populate_table_row(t_appts, idx, row, col_widths=dict_widths, zebra=(idx % 2 == 0))
        
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Bảng MEDICAL_RECORDS
    add_styled_heading(doc, "Bảng 4.2.4: MEDICAL_RECORDS (Hồ sơ bệnh án điện tử)", level=3)
    t_records = create_styled_table(doc, rows=8, cols=4, headers=dict_headers, col_widths=dict_widths)
    rec_data = [
        ("record_id", "INT", "PK, Auto Increment", "Mã hồ sơ bệnh án."),
        ("patient_id", "INT", "FK -> PATIENTS", "Mã bệnh nhân chủ sở hữu bệnh án."),
        ("doctor_id", "INT", "FK -> DOCTORS", "Mã bác sĩ lập bệnh án."),
        ("record_date", "DATE", "NOT NULL", "Ngày lập bệnh án."),
        ("diagnosis", "TEXT", "NOT NULL", "Chẩn đoán y khoa chính."),
        ("prescription", "TEXT", "NOT NULL", "Đơn thuốc chi tiết (tên thuốc, liều dùng, tần suất)."),
        ("note", "TEXT", "Nullable", "Ghi chú bổ sung hoặc lời dặn của bác sĩ.")
    ]
    for idx, row in enumerate(rec_data, start=1):
        populate_table_row(t_records, idx, row, col_widths=dict_widths, zebra=(idx % 2 == 0))
        
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Bảng BILLS
    add_styled_heading(doc, "Bảng 4.2.5: BILLS (Thông tin hóa đơn viện phí)", level=3)
    t_bills = create_styled_table(doc, rows=8, cols=4, headers=dict_headers, col_widths=dict_widths)
    bill_data = [
        ("bill_id", "INT", "PK, Auto Increment", "Mã định danh hóa đơn."),
        ("appointment_id", "INT", "FK -> APPOINTMENTS", "Mã cuộc hẹn liên kết."),
        ("patient_id", "INT", "FK -> PATIENTS", "Mã bệnh nhân nhận hóa đơn."),
        ("total_amount", "DECIMAL(10,2)", "NOT NULL", "Tổng số tiền cần thanh toán."),
        ("status", "VARCHAR(20)", "NOT NULL", "Trạng thái hóa đơn: UNPAID, PAID, CANCELLED."),
        ("payment_method", "VARCHAR(30)", "Nullable", "Phương thức: CASH, CREDIT_CARD, BANK_TRANSFER, E_WALLET."),
        ("bill_date", "DATETIME", "NOT NULL", "Ngày giờ lập hóa đơn y tế.")
    ]
    for idx, row in enumerate(bill_data, start=1):
        populate_table_row(t_bills, idx, row, col_widths=dict_widths, zebra=(idx % 2 == 0))

    doc.add_page_break()

    # --- SECTION 5: YÊU CẦU PHI CHỨC NĂNG ---
    add_styled_heading(doc, "5. YÊU CẦU PHI CHỨC NĂNG", level=1)
    
    add_styled_heading(doc, "5.1. Hiệu năng (Performance Requirements)", level=2)
    add_bullet_point(doc, "Thời gian phản hồi (Response Time): Các tác vụ tra cứu lịch hẹn và hồ sơ bệnh án thông thường phải có thời gian phản hồi dưới 1.5 giây trong điều kiện tải bình thường.")
    add_bullet_point(doc, "Tải đồng thời (Concurrency): Hệ thống phải chịu tải tối thiểu 1,000 người dùng truy cập đồng thời và xử lý 100 giao dịch đặt lịch/giây mà không xảy ra tình trạng nghẽn kết nối.")
    add_bullet_point(doc, "Độ tin cậy (Reliability): Tỷ lệ hoạt động liên tục (Uptime) đạt 99.9% (chỉ chấp nhận tối đa 8.76 giờ ngưng hoạt động ngoài ý muốn mỗi năm).")
    
    add_styled_heading(doc, "5.2. Bảo mật & An toàn thông tin y tế (Security & HIPAA)", level=2)
    add_body_paragraph(doc, 
        "Hồ sơ sức khỏe cá nhân là dữ liệu tối mật. Hệ thống cam kết thực hiện đầy đủ các tiêu chuẩn bảo mật y tế nghiêm ngặt:"
    )
    add_bullet_point(doc, "Kiểm soát truy cập dựa trên vai trò (RBAC): Chỉ bác sĩ được phân công điều trị mới có quyền truy cập và chỉnh sửa bệnh án của bệnh nhân đó. Nhân viên hành chính chỉ được xem thông tin lịch hẹn và viện phí.")
    add_bullet_point(doc, "Mã hóa dữ liệu: Sử dụng mã hóa AES-256 đối với dữ liệu lưu trữ tại cơ sở dữ liệu (Encryption at Rest) và mã hóa TLS 1.3 đối với toàn bộ dữ liệu truyền tải trên mạng (Encryption in Transit).")
    add_bullet_point(doc, "Nhật ký lịch sử (Audit Trail): Mọi thao tác tạo, đọc, cập nhật, xóa (CRUD) trên hồ sơ bệnh án y khoa đều phải được ghi lại tự động vào hệ thống log tập trung, không thể chỉnh sửa, ghi rõ thời gian và danh tính người thực hiện để phục vụ công tác thanh tra.")
    
    add_styled_heading(doc, "5.3. Khả năng mở rộng & Sẵn sàng (Scalability & Availability)", level=2)
    add_bullet_point(doc, "Hệ thống hỗ trợ cơ chế Tự động giãn nở (Auto-scaling) trên hạ tầng Cloud để tự động tăng cường tài nguyên tính toán (CPU, RAM) vào các giờ cao điểm (sáng sớm từ 7:00 - 9:00).")
    add_bullet_point(doc, "Cơ sở dữ liệu hỗ trợ cơ chế Read-Replica để phân tách tải giữa các tác vụ báo cáo/tra cứu (Read) và các tác vụ nghiệp vụ ghi dữ liệu (Write).")
    
    doc.save(DOCX_FILE)
    print(f"Saved {DOCX_FILE}")

def main():
    # Mermaid diagrams source
    usecase_mermaid = """
graph TD
    Patient((Bệnh nhân))
    Doctor((Bác sĩ))
    Admin((Nhân viên Y tế / Admin))

    subgraph PH_LICH_HEN [Phân hệ Đăng ký & Lịch hẹn]
        UC_Book(Đặt lịch khám)
        UC_ViewAppt(Xem lịch khám)
        UC_CancelAppt(Hủy/Đổi lịch hẹn)
        UC_ManageSched(Quản lý lịch khám bác sĩ)
    end

    subgraph PH_BENH_AN [Phân hệ Hồ sơ bệnh án điện tử]
        UC_ViewEHR(Xem hồ sơ bệnh án)
        UC_UpdateEHR(Cập nhật bệnh án & kê đơn)
        UC_ManageEHR(Quản lý thông tin bệnh án)
    end

    subgraph PH_THANH_TOAN [Phân hệ Viện phí & Thanh toán]
        UC_ViewBill(Xem hóa đơn viện phí)
        UC_PayBill(Thanh toán trực tuyến)
        UC_CreateBill(Lập hóa đơn & biên lai)
    end

    Patient --> UC_Book
    Patient --> UC_ViewAppt
    Patient --> UC_CancelAppt
    Patient --> UC_ViewEHR
    Patient --> UC_ViewBill
    Patient --> UC_PayBill

    Doctor --> UC_ViewAppt
    Doctor --> UC_UpdateEHR
    Doctor --> UC_ViewEHR

    Admin --> UC_ManageSched
    Admin --> UC_ManageEHR
    Admin --> UC_CreateBill
"""

    erd_mermaid = """
erDiagram
    PATIENTS ||--o{ APPOINTMENTS : "has"
    DOCTORS ||--o{ APPOINTMENTS : "assigned_to"
    PATIENTS ||--o{ MEDICAL_RECORDS : "owns"
    DOCTORS ||--o{ MEDICAL_RECORDS : "writes"
    APPOINTMENTS ||--o| BILLS : "generates"
    PATIENTS ||--o{ BILLS : "pays"

    PATIENTS {
        int patient_id PK
        string full_name
        string phone_number
        date date_of_birth
        string gender
        string address
        string email
    }

    DOCTORS {
        int doctor_id PK
        string full_name
        string specialization
        string phone_number
        string email
        string schedule_details
    }

    APPOINTMENTS {
        int appointment_id PK
        int patient_id FK
        int doctor_id FK
        datetime appointment_date
        string reason
        string status
    }

    MEDICAL_RECORDS {
        int record_id PK
        int patient_id FK
        int doctor_id FK
        date record_date
        string diagnosis
        string prescription
        string note
    }

    BILLS {
        int bill_id PK
        int appointment_id FK
        int patient_id FK
        decimal total_amount
        string status
        string payment_method
        datetime bill_date
    }
"""

    # Fetch diagrams
    usecase_img = fetch_mermaid_image(usecase_mermaid, "usecase_diagram.png")
    erd_img = fetch_mermaid_image(erd_mermaid, "erd_diagram.png")

    # Generate documents
    generate_prompts_file()
    generate_docx_file(usecase_img, erd_img)
    print("All tasks completed successfully!")

if __name__ == "__main__":
    def set_font(run, size_pt=11, color_rgb=(0, 0, 0), bold=False, italic=False):
        run.font.name = "Arial"
        run.font.size = Pt(size_pt)
        run.font.color.rgb = RGBColor(*color_rgb)
        run.bold = bold
        run.italic = italic
    main()
